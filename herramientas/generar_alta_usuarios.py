"""
Genera, a partir de un Excel de usuarios (APELLIDOS Y NOMBRES | DNI | PERFIL), los dos scripts del
alta estandar PAU -> SSIPE para un ambiente:

  <lote>_1_PAU.sql    se corre en la base del PAU: es el 24_ALTA_ESTANDAR con @asig ya cargado.
                      Asigna el perfil SSIPE a usuarios que YA existen en PAU (el alta de la cuenta
                      va por el portal PAU) y al final devuelve JsonParaScript25.
  <lote>_2_SSIPE.sql  se corre en la base SSIPE: pegar ahi el JsonParaScript25 y ejecutar.
                      Llama a integracion.paHomologarUsuariosPau (script 25 v2): PauUsuario + directorio.

Ambos arrancan en simulacion (@confirmar = 0). Solo se leen las columnas de nombre, documento y
perfil; cualquier otra columna del Excel (correo, celular, observaciones, claves) se ignora.
Los scripts generados contienen DNIs: guardarlos fuera de git (por ejemplo docs/USUARIOS/lotes).

Uso:
  python generar_alta_usuarios.py usuarios.xlsx --salida <carpeta> --revisado-por <DNI de quien homologa>
      [--ambiente DESA|PROD] [--pau-db X --ssipe-db Y --sistema-id N] [--id-area-defecto N] [--lote nombre]
"""
import argparse
import datetime as dt
import re
import sys
import unicodedata
from pathlib import Path

from openpyxl import load_workbook

AQUI = Path(__file__).resolve().parent
PLANTILLA_PAU = AQUI.parent / "24_ALTA_ESTANDAR_asignar_perfil_PAU.sql"

AMBIENTES = {
    "DESA": {"pau_db": "PVDPAU_PROD", "ssipe_db": "DBSSIPE2", "sistema_id": 2020},
    "PROD": {"pau_db": None, "ssipe_db": None, "sistema_id": None},
}

# Nombre de perfil (como lo escribe el area usuaria) -> CodigoSSO homologado en PauPerfil
PERFILES = {
    "ADMINISTRADOR": "P0001",
    "GERENTE DE OBRA": "P0021",
    "JEFE DE OBRA": "P0022",
    "ADMINISTRADOR DE CONTRATO": "P0023",
    "ADMINISTRADOR DE CONTRATO DE OBRA": "P0023",
    "SUPERVISOR DE OBRA": "P0024",
    "COORDINADOR DE OBRA": "P0025",
    "COORDINADOR EXPEDIENTE": "P0027",
    "COORDINADOR DE EXPEDIENTE": "P0027",
    "ESPECIALISTA DE EXPEDIENTE": "P0028",
    "ESPECIALISTA PREINVERSION": "P0029",
    "ESPECIALISTA DE PREINVERSION": "P0029",
    "RESPONSABLE EXPEDIENTE TECNICO": "P0041",
    "RESPONSABLE PACRI": "P0042",
    "RESPONSABLE EJECUCION": "P0043",
    "COORDINADOR PATS": "P0044",
    "LECTOR GENERAL": "P0045",
    "RESPONSABLE ABASTECIMIENTO": "P0046",
}


def normalizar(texto):
    texto = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", texto).strip().upper()


def sql(texto):
    return "N'" + str(texto).replace("'", "''") + "'"


def leer_excel(ruta):
    ws = load_workbook(ruta, data_only=True).worksheets[0]
    filas = list(ws.iter_rows(values_only=True))
    cab = [normalizar(c) for c in filas[0]]

    def columna(*opciones):
        for i, c in enumerate(cab):
            if any(o in c for o in opciones):
                return i
        sys.exit(f"No se encontro la columna {opciones} en la cabecera: {filas[0]}")

    c_nom, c_doc, c_per = columna("APELLIDOS", "NOMBRE"), columna("DNI", "DOCUMENTO"), columna("PERFIL")
    usuarios, errores = [], []
    for n, fila in enumerate(filas[1:], start=2):
        doc = re.sub(r"\D", "", str(fila[c_doc] or "")) if fila[c_doc] is not None else ""
        if not doc and not fila[c_nom]:
            continue
        perfil_txt = normalizar(fila[c_per])
        codigo = perfil_txt if re.fullmatch(r"P\d{4}", perfil_txt) else PERFILES.get(perfil_txt)
        if not doc:
            errores.append(f"fila {n}: sin DNI/documento")
        elif not codigo:
            errores.append(f"fila {n}: perfil '{fila[c_per]}' no reconocido (agregarlo en PERFILES o usar el codigo Pxxxx)")
        else:
            usuarios.append((doc.zfill(8), normalizar(fila[c_nom]), codigo))
    if errores:
        sys.exit("Excel con errores:\n  " + "\n  ".join(errores))
    return usuarios


def script_pau(usuarios, amb):
    texto = PLANTILLA_PAU.read_text(encoding="utf-8")
    ini = texto.index("-- =====> UNICO BLOQUE A EDITAR")
    fin = texto.index("-- =====> FIN DEL BLOQUE A EDITAR <=====")
    filas = ",\n".join(f" ({sql(d)}, {sql(n)}, '{c}')" for d, n, c in usuarios)
    bloque = ("-- =====> UNICO BLOQUE A EDITAR EN CADA LOTE NUEVO <=====  (generado desde Excel)\n"
              "DECLARE @asig TABLE (Documento varchar(20), NombreCompleto nvarchar(200), CodigoSSO varchar(10));\n"
              f"INSERT @asig VALUES\n{filas};\n")
    texto = texto[:ini] + bloque + texto[fin:]
    texto = texto.replace("IF DB_NAME() <> N'PVDPAU_PROD'", f"IF DB_NAME() <> N'{amb['pau_db']}'", 1)
    texto = re.sub(r"DECLARE @sis int = \d+;", f"DECLARE @sis int = {amb['sistema_id']};", texto, count=1)
    return texto


def script_ssipe(usuarios, amb, revisado_por, id_area):
    docs = ", ".join(sql(d) for d, _, _ in usuarios)
    area = "NULL   -- <<< IdArea SSIPE para usuarios nuevos (ver EXEC integracion.paListarArea)" if id_area is None else str(id_area)
    return f"""/* Alta estandar PAU -> SSIPE, FASE 2: homologar en la base SSIPE ({len(usuarios)} usuarios).
   1) Correr antes el _1_PAU.sql de este lote con @confirmar = 1.
   2) Pegar abajo el valor de JsonParaScript25 que devolvio.
   3) Ejecutar con @confirmar = 0, revisar Estado/Directorio por fila, y repetir con @confirmar = 1. */
IF DB_NAME() <> N'{amb['ssipe_db']}' OR OBJECT_ID(N'integracion.paHomologarUsuariosPau', N'P') IS NULL
BEGIN RAISERROR(N'Base incorrecta o falta integracion.paHomologarUsuariosPau (script 25): ejecucion cancelada.', 16, 1); SET NOEXEC ON; END
GO
DECLARE @confirmar bit = 0;
DECLARE @json nvarchar(max) = N'<<< PEGAR AQUI JsonParaScript25 >>>';
IF ISJSON(@json) <> 1 THROW 58001, 'Pegar en @json el JsonParaScript25 del script PAU del lote.', 1;

EXEC integracion.paHomologarUsuariosPau
    @usuariosJson = @json,
    @sistemaId = {amb['sistema_id']},
    @revisadoPor = {sql(revisado_por)},
    @idAreaPorDefecto = {area},
    @confirmar = @confirmar;

SELECT Verificacion = 'Candidatos visibles en Asignar Proyecto / sesion', IdUsuario, Documento, CodigoPerfil, NombrePerfil, Area
FROM integracion.vw_UsuarioSsipe WHERE Documento IN ({docs});
GO
SET NOEXEC OFF;
GO
"""


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("excel")
    ap.add_argument("--salida", required=True)
    ap.add_argument("--revisado-por", required=True, help="documento de quien homologa (auditoria)")
    ap.add_argument("--ambiente", choices=AMBIENTES, default="DESA")
    ap.add_argument("--pau-db")
    ap.add_argument("--ssipe-db")
    ap.add_argument("--sistema-id", type=int)
    ap.add_argument("--id-area-defecto", type=int)
    ap.add_argument("--lote", default=dt.date.today().isoformat())
    a = ap.parse_args()

    amb = dict(AMBIENTES[a.ambiente])
    for clave in ("pau_db", "ssipe_db", "sistema_id"):
        valor = getattr(a, clave)
        if valor is not None:
            amb[clave] = valor
        if amb[clave] is None:
            sys.exit(f"Ambiente {a.ambiente}: indicar --{clave.replace('_', '-')}")

    usuarios = leer_excel(a.excel)
    salida = Path(a.salida)
    salida.mkdir(parents=True, exist_ok=True)
    base = f"{a.lote}_{a.ambiente}"
    (salida / f"{base}_1_PAU.sql").write_text(script_pau(usuarios, amb), encoding="utf-8")
    (salida / f"{base}_2_SSIPE.sql").write_text(script_ssipe(usuarios, amb, a.revisado_por, a.id_area_defecto), encoding="utf-8")
    print(f"{len(usuarios)} usuarios -> {salida / (base + '_1_PAU.sql')} y {salida / (base + '_2_SSIPE.sql')}")
    for d, n, c in usuarios:
        print(f"  {c}  {n}")


if __name__ == "__main__":
    main()
